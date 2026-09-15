#!/usr/bin/env python3
"""Structural validator for Paper Record outputs.

This checks file/layout invariants only; it does not certify the truth of paper
content. The human acceptance gate remains mandatory.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Dict, List

REQUIRED_DIRS = [
    "00-系统说明",
    "01-精华论文",
    "02-研究主题",
    "99-佐证资料/完整论文笔记",
    "99-佐证资料/图表与证据卡",
    "03-研究空白与可复用资产/概念与术语",
    "03-研究空白与可复用资产/研究空白",
    "03-研究空白与可复用资产/可复用科研资产",
    "99-佐证资料/参考文献",
    "09-模板",
    "99-佐证资料/原始论文",
    "99-待人工复核",
]


def rel(path: Path, root: Path) -> str:
    return path.relative_to(root).as_posix()




def frontmatter_keys(text: str) -> set[str]:
    """Return simple YAML frontmatter keys without requiring PyYAML."""
    if not text.startswith("---"):
        return set()
    parts = text.split("---", 2)
    if len(parts) < 3:
        return set()
    return {
        match.group(1)
        for line in parts[1].splitlines()
        if (match := re.match(r"^([A-Za-z_][A-Za-z0-9_-]*):", line.strip()))
    }


def embedded_targets(text: str) -> list[str]:
    return [
        match.group(1).split("|", 1)[0].strip()
        for match in re.finditer(r"!\[\[([^\]]+)\]\]", text)
    ]


def target_exists(target: str, note: Path, root: Path) -> bool:
    target_path = Path(target)
    if target_path.is_absolute():
        return target_path.exists()
    return (note.parent / target_path).exists() or (root / target_path).exists() or any(
        candidate.name == target_path.name for candidate in root.rglob(target_path.name)
    )

def check(condition: bool, name: str, detail: str, errors: List[dict], warnings: List[dict]) -> None:
    item = {"check": name, "detail": detail, "status": "pass" if condition else "fail"}
    (warnings if name.startswith("warning:") else errors).append(item) if not condition else None


def main() -> int:
    parser = argparse.ArgumentParser(description="验证Paper Record工作区结构")
    parser.add_argument("--root", required=True, help="Vault或工作目录")
    parser.add_argument("--stage", type=int, choices=range(0, 8), default=0)
    args = parser.parse_args()

    root = Path(args.root).expanduser().resolve()
    errors: List[dict] = []
    warnings: List[dict] = []
    info: Dict[str, object] = {"root": str(root), "stage": args.stage}

    check(root.exists(), "root_exists", "目标目录存在", errors, warnings)
    if not root.exists():
        print(json.dumps({"passed": False, "errors": errors, "warnings": warnings}, ensure_ascii=False, indent=2))
        return 1

    missing_dirs = [d for d in REQUIRED_DIRS if not (root / d).is_dir()]
    check(not missing_dirs, "required_directories", "缺失：" + ", ".join(missing_dirs) if missing_dirs else "全部目录存在", errors, warnings)

    manifest = root / "00-系统说明" / "文件分类清单.md"
    check(manifest.exists(), "file_classification_manifest", "文件分类清单存在", errors, warnings)
    manifest_text = manifest.read_text(encoding="utf-8") if manifest.exists() else ""

    all_files = [p for p in root.rglob("*") if p.is_file() and ".git" not in p.parts]
    info["file_count"] = len(all_files)
    unregistered = []
    for p in all_files:
        r = rel(p, root)
        if r == "00-系统说明/文件分类清单.md":
            continue
        # Files under recognized business directories are expected to be registered.
        if any(r.startswith(d + "/") for d in REQUIRED_DIRS) and r not in manifest_text:
            unregistered.append(r)
    if unregistered:
        warnings.append({"check": "warning:unregistered_files", "detail": "未在分类清单中出现：" + "; ".join(unregistered[:20]), "status": "warning"})

    if args.stage == 0:
        templates = [
            "09-模板/论文笔记模板.md",
            "09-模板/图表证据卡模板.md",
            "09-模板/分类小结模板.md",
        ]
        missing_templates = [t for t in templates if not (root / t).exists()]
        check(not missing_templates, "starter_templates", "缺失：" + ", ".join(missing_templates) if missing_templates else "基础模板存在", errors, warnings)

    if args.stage == 1:
        index_files = list((root / "01-精华论文").glob("*.md")) if (root / "01-精华论文").exists() else []

        def has_real_rows(path: Path) -> bool:
            text = path.read_text(encoding="utf-8", errors="replace")
            rows = [line.strip() for line in text.splitlines() if line.strip().startswith("|")]
            # A header plus separator is an empty starter; a third table row is
            # the minimum evidence that an imported paper is present. The
            # starter marker is intentionally retained after real rows are
            # appended, so it must not invalidate an otherwise real index.
            return len(rows) >= 3

        real_indexes = [p for p in index_files if has_real_rows(p)]
        info["index_file_count"] = len(index_files)
        info["real_index_file_count"] = len(real_indexes)
        check(
            bool(real_indexes),
            "stage1_index",
            "已发现包含论文条目的分类索引文件"
            if real_indexes
            else "仅发现初始化索引模板，不能判定 Stage 1 完成",
            errors,
            warnings,
        )

    if args.stage == 2:
        missing_csv = list((root / "99-待人工复核").glob("无法下载全文*.csv")) if (root / "99-待人工复核").exists() else []
        attachment_files = list((root / "99-佐证资料/原始论文").rglob("*.pdf")) if (root / "99-佐证资料/原始论文").exists() else []
        info["pdf_count"] = len(attachment_files)
        invalid_pdfs = []
        for pdf in attachment_files:
            try:
                if pdf.read_bytes()[:5] != b"%PDF-":
                    invalid_pdfs.append(rel(pdf, root))
            except OSError:
                invalid_pdfs.append(rel(pdf, root) + " (无法读取)")
        if invalid_pdfs:
            errors.append({"check": "pdf_signature", "detail": "文件扩展名为PDF但文件头无 `%PDF-`：" + "; ".join(invalid_pdfs[:20]), "status": "fail"})
        if not attachment_files and not missing_csv:
            warnings.append({"check": "warning:stage2_no_pdf_or_failure_csv", "detail": "没有发现PDF，也没有无法下载清单；需人工确认阶段是否实际执行", "status": "warning"})

    if args.stage == 3:
        review_files = list((root / "99-待人工复核").rglob("*")) if (root / "99-待人工复核").exists() else []
        info["manual_review_file_count"] = len([p for p in review_files if p.is_file()])

    if args.stage == 4:
        excel_files = list((root / "01-精华论文").glob("*.xlsx")) if (root / "01-精华论文").exists() else []
        check(bool(excel_files), "stage4_excel", "已发现Excel索引文件" if excel_files else "未发现Excel索引文件", errors, warnings)
        if excel_files:
            try:
                from openpyxl import load_workbook
            except ImportError:
                warnings.append({"check": "warning:excel_not_tested", "detail": "未安装 openpyxl，无法验证工作簿、Sheet 和表头", "status": "warning"})
            else:
                invalid_workbooks = []
                empty_sheets = []
                for workbook_path in excel_files:
                    try:
                        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
                        if not workbook.sheetnames:
                            invalid_workbooks.append(rel(workbook_path, root) + " (无Sheet)")
                        for sheet in workbook.worksheets:
                            if sheet.max_row < 2 or sheet.max_column < 1:
                                empty_sheets.append(rel(workbook_path, root) + "!" + sheet.title)
                        workbook.close()
                    except Exception as exc:  # openpyxl raises several format-specific errors
                        invalid_workbooks.append(rel(workbook_path, root) + " (" + type(exc).__name__ + ")")
                if invalid_workbooks:
                    errors.append({"check": "excel_workbook", "detail": "Excel无法读取或缺少Sheet：" + "; ".join(invalid_workbooks[:20]), "status": "fail"})
                if empty_sheets:
                    warnings.append({"check": "warning:excel_empty_sheet", "detail": "发现空Sheet：" + "; ".join(empty_sheets[:20]), "status": "warning"})

    if args.stage == 5:
        tag_dict = list((root / "01-精华论文").glob("*标签*.xlsx")) + list((root / "00-系统说明").glob("*标签*.md"))
        if not tag_dict:
            warnings.append({"check": "warning:stage5_no_tag_dictionary", "detail": "未发现标签字典；可在Zotero中完成标签后再补充", "status": "warning"})

    if args.stage == 6:
        bib_files = list((root / "99-佐证资料/参考文献").glob("*.bib"))
        ref_excels = list((root / "99-佐证资料/参考文献").glob("*.xlsx"))
        ref_files = bib_files + ref_excels
        check(bool(ref_files), "stage6_references", "已发现BibTeX/Excel参考文献文件" if ref_files else "未发现参考文献文件", errors, warnings)
        bad_bib = []
        for bib in bib_files:
            text = bib.read_text(encoding="utf-8", errors="replace")
            entries = re.findall(r"(?m)^\s*@([A-Za-z]+)\s*\{", text)
            if not entries:
                bad_bib.append(rel(bib, root) + " (没有BibTeX条目)")
            if text.count("{") != text.count("}"):
                bad_bib.append(rel(bib, root) + " (花括号不平衡)")
        if bad_bib:
            errors.append({"check": "bibtex_syntax", "detail": "BibTeX结构异常：" + "; ".join(bad_bib[:20]), "status": "fail"})
        if ref_excels:
            try:
                from openpyxl import load_workbook
            except ImportError:
                warnings.append({"check": "warning:reference_excel_not_tested", "detail": "未安装 openpyxl，无法验证引用Excel的master/occurrence工作表", "status": "warning"})
            else:
                missing_relation_sheets = []
                for workbook_path in ref_excels:
                    try:
                        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
                        names = {name.lower() for name in workbook.sheetnames}
                        if not any("master" in name or "总表" in name for name in names):
                            missing_relation_sheets.append(rel(workbook_path, root) + " (缺master/总表)")
                        if not any("occurrence" in name or "溯源" in name or "来源" in name for name in names):
                            missing_relation_sheets.append(rel(workbook_path, root) + " (缺occurrence/溯源/来源)")
                        workbook.close()
                    except Exception as exc:
                        missing_relation_sheets.append(rel(workbook_path, root) + " (" + type(exc).__name__ + ")")
                if missing_relation_sheets:
                    errors.append({"check": "reference_excel_relations", "detail": "引用Excel关系表不完整：" + "; ".join(missing_relation_sheets[:20]), "status": "fail"})

    if args.stage == 7:
        paper_notes = list((root / "99-佐证资料/完整论文笔记").rglob("*.md")) if (root / "99-佐证资料/完整论文笔记").exists() else []
        figure_cards = list((root / "99-佐证资料/图表与证据卡").rglob("*.md")) if (root / "99-佐证资料/图表与证据卡").exists() else []
        images = [p for p in (root / "99-佐证资料/图表与证据卡").rglob("*") if p.suffix.lower() in {".png", ".jpg", ".jpeg", ".svg", ".pdf"}] if (root / "99-佐证资料/图表与证据卡").exists() else []
        info.update({"paper_note_count": len(paper_notes), "figure_card_count": len(figure_cards), "figure_asset_count": len(images)})
        check(bool(paper_notes), "stage7_paper_notes", "已发现论文笔记" if paper_notes else "未发现论文笔记", errors, warnings)
        if not figure_cards and images:
            warnings.append({"check": "warning:figure_without_card", "detail": "发现图表素材但没有图表证据卡", "status": "warning"})

        # Validate the minimum machine-readable contract, not only file presence.
        required_paper_keys = {
            "title", "zotero_item_id", "category", "reading_status",
            "fulltext_status", "evidence_status", "record_batch",
        }
        bad_notes = []
        broken_links = []
        for p in paper_notes:
            text = p.read_text(encoding="utf-8", errors="replace")
            keys = frontmatter_keys(text)
            if not text.startswith("---") or "研究问题" not in text or "论文论证链" not in text:
                bad_notes.append(rel(p, root))
            missing = sorted(required_paper_keys - keys)
            if missing:
                bad_notes.append(rel(p, root) + " (缺少: " + ", ".join(missing) + ")")
            for target in embedded_targets(text):
                if not target_exists(target, p, root):
                    broken_links.append(rel(p, root) + " -> " + target)
        if bad_notes:
            errors.append({"check": "paper_note_contract", "detail": "论文笔记结构或frontmatter不完整：" + "; ".join(bad_notes[:20]), "status": "fail"})
        if broken_links:
            errors.append({"check": "paper_note_embeds", "detail": "图表/附件链接不存在：" + "; ".join(broken_links[:20]), "status": "fail"})

        bad_cards = []
        for p in figure_cards:
            text = p.read_text(encoding="utf-8", errors="replace")
            keys = frontmatter_keys(text)
            required_card_keys = {"record_type", "paper_item_id", "figure_id", "source_page", "verification_status", "record_batch"}
            missing = sorted(required_card_keys - keys)
            if missing:
                bad_cards.append(rel(p, root) + " (缺少: " + ", ".join(missing) + ")")
            for target in embedded_targets(text):
                if not target_exists(target, p, root):
                    broken_links.append(rel(p, root) + " -> " + target)
        if bad_cards:
            errors.append({"check": "figure_card_contract", "detail": "图表证据卡frontmatter不完整：" + "; ".join(bad_cards[:20]), "status": "fail"})

    passed = not errors
    result = {
        "passed": passed,
        "human_acceptance_required": True,
        "errors": errors,
        "warnings": warnings,
        "info": info,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
